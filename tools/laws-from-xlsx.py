#!/usr/bin/env python3
"""
Makes the Old English laws database: one TEI XML file per clause row of the
"OE Laws Database" spreadsheet.

    python3 tools/laws-from-xlsx.py OE_Laws_Database_For_Claude.xlsx
    python3 tools/laws-from-xlsx.py OE_Laws_Database_For_Claude.xlsx --out oe-laws

Needs Python 3 and openpyxl (pip install openpyxl). Writes the files into
oe-laws/ (or --out), plus oe-laws/index.json, a catalogue of every clause.
Files from an earlier run are replaced. See docs/OE-LAWS.md for what each
column of the spreadsheet becomes.

How the spreadsheet is read
- Clause rows have a law code abbreviation in column B. Each becomes one file.
- Sub-header rows (column E starts "Place and Date of Promulgation") repeat
  before each block of clauses. They are not clauses: they say which law,
  which editor and which manuscripts the rows below them use, and that goes
  into the header of every file in the block.
- Empty rows are skipped.
- Where the spreadsheet merges cells over several clause rows, the merged text
  is copied to each of those rows and the file says so.
- In the MSS column (F) the editor's selected manuscript is in bold; that bold
  is kept (witness rend="bold").
"""
import argparse
import datetime
import json
import os
import re
import sys
import unicodedata
from xml.sax.saxutils import escape

import openpyxl
from openpyxl.cell.rich_text import CellRichText, TextBlock

PUBLISHER = 'eLiebermann: An Online Dictionary of Pre-Conquest English Law'
SOURCE_NAME = 'OE Laws Database spreadsheet'

# Same folding as ELieb.slugify() in scripts/eliebermann.js, so file names
# follow the conventions of the dictionary entries.
FOLD = {'æ': 'ae', 'Æ': 'Ae', 'þ': 'th', 'Þ': 'Th', 'ð': 'd', 'Ð': 'D', 'ƿ': 'w', 'Ƿ': 'W',
        'ȝ': 'g', 'Ȝ': 'G', 'ß': 'ss', 'œ': 'oe', 'Œ': 'Oe', '⁊': 'and', 'ꝥ': 'thaet', 'Ꝥ': 'Thaet'}

# Spreadsheet abbreviations that are not spelled exactly like the law code list
# (data/markLaw.json) but clearly mean one of its codes. Checked by eye: see
# docs/OE-LAWS.md. Anything not matched either way gets no law code link.
ASSUMED_ANCHORS = {
    'ord': 'Ordal', 'becw': 'Becwaeth', 'mirc': 'Mirce', 'norðleod': 'Northleoda',
    'norcyricgrið': 'NorGrith', 'rom': 'Romscot',
    'promreg': 'SacrCor',       # the dictionary cites the Promissio regis as "Sacr cor"
    'wllad': 'WiLad', 'hncom': 'HnComm',
}

ILLEGAL_XML = re.compile('[\x00-\x08\x0b\x0c\x0e-\x1f￾￿]')


def clean(value):
    """Cell text with every run of spaces (including non-breaking ones) as one space."""
    if value is None:
        return ''
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = ILLEGAL_XML.sub('', str(value))
    return re.sub(r'\s+', ' ', text.replace('\xa0', ' ')).strip()


def slugify(text):
    s = ''.join(FOLD.get(c, c) for c in str(text))
    s = unicodedata.normalize('NFD', s)
    s = ''.join(c for c in s if not unicodedata.combining(c))
    return re.sub(r'[^A-Za-z0-9]+', '-', s).strip('-')


def code_key(code):
    """A law code abbreviation with spaces, full stops and colons ignored."""
    return re.sub(r'[\s.:]+', '', unicodedata.normalize('NFC', code)).lower()


LONG_CLAUSE = 60      # a "clause number" longer than this is a note, not a number
MAX_SLUG = 40         # longest clause part of a file name


def clause_key(clause):
    """Clause number in the form dictionary citations can be matched against:
    "23,2" and "23.2" agree, and brackets and capitals are ignored."""
    return re.sub(r'[\s\[\]]+', '', clause.replace(',', '.')).lower()


def clause_keys(clause):
    """One key per clause number. Most rows have one number; a few cover
    several ("3, 3.1, 3.2"), and then each of them is a key. A comma with no
    space after it is Liebermann's decimal point ("23,2"), not a separator."""
    return [clause_key(p) for p in clause_parts(clause)]


def clause_parts(clause):
    if not clause:
        return []
    return [p for p in re.split(r'\s*,\s+|\.\s+(?=\d)', clause) if p.strip()] or [clause]


def esc(text):
    return escape(text)


def attr(text):
    return escape(text, {'"': '&quot;'})


# ----------------------------------------------------------------------
# Reading the spreadsheet
# ----------------------------------------------------------------------

def load(path):
    values = openpyxl.load_workbook(path, data_only=True).active      # formulas as their results
    rich = openpyxl.load_workbook(path, rich_text=True).active         # bold runs
    merged = {}
    for rng in values.merged_cells.ranges:
        top = values.cell(rng.min_row, rng.min_col).value
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                if (r, c) != (rng.min_row, rng.min_col):
                    merged[(r, c)] = (clean(top), rng.min_row, rng.max_row)
    return values, rich, merged


def bold_flags(cell):
    """The cell's text as [(character, is_bold)], spaces tidied."""
    base_bold = bool(cell.font and cell.font.b)
    value = cell.value
    parts = []
    if isinstance(value, CellRichText):
        for part in value:
            if isinstance(part, TextBlock):
                parts.append((part.text, bool(part.font.b) if part.font and part.font.b is not None else base_bold))
            else:
                parts.append((str(part), base_bold))
    else:
        parts.append(('' if value is None else str(value), base_bold))
    chars = []
    for text, bold in parts:
        for ch in ILLEGAL_XML.sub('', text).replace('\xa0', ' '):
            ch = ' ' if ch.isspace() else ch
            if ch == ' ' and (not chars or chars[-1][0] == ' '):
                continue
            chars.append((ch, bold))
    while chars and chars[-1][0] == ' ':
        chars.pop()
    return chars


def witnesses(cell):
    """Split the MSS cell into [(text, selected)] at commas and semicolons
    outside brackets. A witness is selected if any of its letters is bold."""
    chars = bold_flags(cell)
    out, current, depth = [], [], 0
    def flush():
        text = ''.join(c for c, _ in current).strip()
        if text:
            out.append((text, any(b for c, b in current if c.isalnum())))
        current.clear()
    for ch, bold in chars:
        if ch in '([':
            depth += 1
        elif ch in ')]':
            depth = max(0, depth - 1)
        if ch in ',;' and depth == 0:
            flush()
        else:
            current.append((ch, bold))
    flush()
    return out


# ----------------------------------------------------------------------
# Matching law codes to the site's law code list
# ----------------------------------------------------------------------

def anchor_finder(mark_law_path):
    by_text = {}
    if mark_law_path and os.path.exists(mark_law_path):
        for row in json.load(open(mark_law_path, encoding='utf-8')):
            by_text.setdefault(code_key(row['Gesetze Text']), set()).add(row['Anchor'])
    def find(code):
        key = code_key(code)
        found = by_text.get(key)
        if found and len(found) == 1:
            return next(iter(found)), 'matched'
        if key in ASSUMED_ANCHORS and ASSUMED_ANCHORS[key] in {a for s in by_text.values() for a in s}:
            return ASSUMED_ANCHORS[key], 'assumed'
        return '', 'none'
    return find


# ----------------------------------------------------------------------
# Building one file
# ----------------------------------------------------------------------

PREFIX_DEFS = '''      <listPrefixDef>
        <prefixDef ident="law" matchPattern="(.+)" replacementPattern="lawAbb.html#$1"/>
        <prefixDef ident="gram" matchPattern="(.+)" replacementPattern="gramAbb.html#$1"/>
        <prefixDef ident="entry" matchPattern="(.+)" replacementPattern="entry.html?id=$1"/>
      </listPrefixDef>'''


def split_terms(text):
    return [t.strip() for t in text.split(';') if t.strip()]


def after_colon(label):
    """'Translation of H: Oliver' -> 'Oliver'; no colon -> ''."""
    return label.rsplit(':', 1)[1].strip() if ':' in label else ''


def build(row, block, anchor, today):
    a, code, clause = row['citation'], row['code'], row['clause']
    L = []
    w = L.append
    w('<?xml version="1.0" encoding="UTF-8"?>')
    w('<TEI xmlns="http://www.tei-c.org/ns/1.0" xml:id="%s">' % attr(row['id']))
    w('  <teiHeader>')
    w('    <fileDesc>')
    w('      <titleStmt>')
    w('        <title>%s</title>' % esc(a))
    editor = block.get('editor', '')
    translator = after_colon(block.get('translationLabel', '')) or editor
    if editor:
        w('        <respStmt><resp>Edition of the selected text</resp><name>%s</name></respStmt>' % esc(editor))
    if translator and row['translation']:
        w('        <respStmt><resp>Translation of the selected text</resp><name>%s</name></respStmt>' % esc(translator))
    w('      </titleStmt>')
    w('      <publicationStmt>')
    w('        <publisher>%s</publisher>' % esc(PUBLISHER))
    w('        <idno type="slug">%s</idno>' % esc(row['id']))
    w('        <idno type="citation">%s</idno>' % esc(a))
    w('        <idno type="lawCode">%s</idno>' % esc(code))
    if anchor:
        w('        <idno type="lawCodeAnchor">%s</idno>' % esc(anchor))
    if clause:
        w('        <idno type="clause">%s</idno>' % esc(clause))
        for key in clause_keys(clause):
            w('        <idno type="citationKey">%s</idno>' % esc((anchor or code) + ' ' + key))
    if row['legalText'] and row['legalText'] != code:
        w('        <idno type="legalText">%s</idno>' % esc(row['legalText']))
    w('      </publicationStmt>')
    w('      <sourceDesc>')
    w('        <bibl type="edition">')
    w('          <author>Felix Liebermann</author>, <title>Die Gesetze der Angelsachsen</title>, <biblScope unit="volume">1</biblScope>%s' %
      (', <biblScope unit="page">%s</biblScope>' % esc(row['page']) if row['page'] else ''))
    w('        </bibl>')
    if block.get('title'):
        reign = block.get('reign', '')
        w('        <bibl type="legalText"><title>%s</title>%s</bibl>' %
          (esc(block['title']), (' <note type="reign">%s</note>' % esc(reign)) if reign else ''))
    if row['witnesses']:
        w('        <listWit>')
        for text, selected in row['witnesses']:
            w('          <witness%s>%s</witness>' % (' rend="bold"' if selected else '', esc(text)))
        w('        </listWit>')
    if block.get('selectedLabel'):
        w('        <bibl type="selectedText">%s</bibl>' % esc(block['selectedLabel']))
    supp = block.get('supplementaryLabel', '')
    if row['supplementary'] and supp and not supp.startswith('Supplementary MS'):
        w('        <bibl type="supplementaryText">%s</bibl>' % esc(supp))
    if row['liebermannNumber'] or row['editorNumber']:
        w('        <bibl type="numbering">')
        if row['liebermannNumber']:
            w('          <idno type="liebermannNumber">%s</idno>' % esc(row['liebermannNumber']))
        if row['editorNumber']:
            w('          <idno type="editorNumber"%s>%s</idno>' %
              ((' n="%s"' % attr(editor)) if editor else '', esc(row['editorNumber'])))
        w('        </bibl>')
    w('      </sourceDesc>')
    w('    </fileDesc>')
    w('    <encodingDesc>')
    w(PREFIX_DEFS)
    w('    </encodingDesc>')
    keyword_lists = [('offence-legal-event', row['offence']), ('people', row['people']), ('locations', row['locations'])]
    if row['promulgation'] or any(k for _, k in keyword_lists):
        w('    <profileDesc>')
        if row['promulgation']:
            w('      <creation><note type="promulgation">%s</note></creation>' % esc(row['promulgation']))
        if any(k for _, k in keyword_lists):
            w('      <textClass>')
            for name, terms in keyword_lists:
                if terms:
                    w('        <keywords n="%s">%s</keywords>' % (name, ''.join('<term>%s</term>' % esc(t) for t in split_terms(terms))))
            w('      </textClass>')
        w('    </profileDesc>')
    w('    <revisionDesc>')
    w('      <change when="%s">Generated from row %d of the %s; not yet checked against the print.</change>' % (today, row['sheetRow'], SOURCE_NAME))
    w('    </revisionDesc>')
    w('  </teiHeader>')
    w('  <text>')
    w('    <body>')
    w('      <div type="clause"%s>' % ((' n="%s"' % attr(clause)) if clause else ''))
    head_code = '<ref type="lawCode" target="law:%s">%s</ref>' % (attr(anchor), esc(code)) if anchor else esc(code)
    w('        <head>%s%s</head>' % (head_code, (' ' + esc(clause)) if clause else ''))
    texts = [('selectedText', 'und', row['selected']), ('translation', 'en', row['translation']),
             ('secondTranslation', 'en', row['secondTranslation']),
             ('supplementaryText', 'und', row['supplementary']), ('supplementaryTranslation', 'en', row['supplementaryTranslation'])]
    for kind, lang, text in texts:
        if text:
            w('        <ab type="%s" xml:lang="%s">%s</ab>' % (kind, lang, esc(text)))
    if row.get('clauseNote'):
        w('        <note type="clause">%s</note>' % esc(row['clauseNote']))
    for note in row['mergedNotes']:
        w('        <note type="source">%s</note>' % esc(note))
    w('      </div>')
    w('    </body>')
    w('  </text>')
    w('</TEI>')
    return '\n'.join(L) + '\n'


# ----------------------------------------------------------------------
# Main
# ----------------------------------------------------------------------

def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('xlsx')
    ap.add_argument('--out', default='oe-laws')
    ap.add_argument('--mark-law', default='data/markLaw.json', help='law code list used to find each code\'s Anchor')
    ap.add_argument('--date', default=datetime.date.today().isoformat())
    args = ap.parse_args()

    values, rich, merged = load(args.xlsx)
    find_anchor = anchor_finder(args.mark_law)

    def cell(r, c):
        """Cleaned text of a cell, taking the merged cell's text if it is part of one."""
        v = clean(values.cell(r, c).value)
        if v or (r, c) not in merged:
            return v, None
        text, first, last = merged[(r, c)]
        return text, (first, last)

    COL = {k: i for i, k in enumerate('ABCDEFGHIJKLMNOPQ', start=1)}
    block = {}
    rows, used, used_lower, part_counter = [], set(), set(), {}
    skipped_empty = headers = 0

    for r in range(2, values.max_row + 1):
        get = lambda letter: cell(r, COL[letter])[0]
        if get('E').startswith('Place and Date'):
            headers += 1
            title = get('D')
            m = re.match(r'^(.*?)\s+(r\.\s*.*)$', title)
            block = {'title': m.group(1) if m else title, 'reign': m.group(2) if m else '',
                     'editor': get('I'), 'selectedLabel': get('J'), 'translationLabel': get('K'),
                     'secondLabel': get('L'), 'supplementaryLabel': get('M'), 'sheetRow': r}
            continue
        code = get('B')
        if not code:
            skipped_empty += 1
            continue
        clause = get('C')
        clause_note = ''
        if len(clause) > LONG_CLAUSE:           # prose typed into the clause column
            clause, clause_note = '', clause
        anchor, how = find_anchor(code)
        merged_notes = []
        def text(letter):
            v, span = cell(r, COL[letter])
            if span and v:
                merged_notes.append('The text in the spreadsheet column "%s" is one merged cell shared by rows %d to %d, so the same text appears in each of those clauses.'
                                    % (letter, span[0], span[1]))
            return v
        row = {
            'sheetRow': r, 'code': code, 'clause': clause, 'anchor': anchor, 'anchorHow': how,
            'clauseNote': clause_note,
            'citation': (get('A') or ('%s %s' % (code, clause)).strip()) if not clause_note else code,
            'legalText': get('D'), 'promulgation': get('E'), 'page': get('G'),
            'liebermannNumber': text('H'), 'editorNumber': text('I'),
            'selected': text('J'), 'translation': text('K'), 'secondTranslation': text('L'),
            'supplementary': text('M'), 'supplementaryTranslation': text('N'),
            'offence': text('O'), 'people': text('P'), 'locations': text('Q'),
            'witnesses': witnesses(rich.cell(r, COL['F'])), 'mergedNotes': sorted(set(merged_notes)),
            'block': dict(block),
        }
        # file name and xml:id: law code (its Anchor if there is one) + clause number
        base = anchor or slugify(code) or 'code'
        if clause:
            parts = clause_parts(clause)
            part = slugify(parts[0] if len(parts) == 1 else parts[0] + ' to ' + parts[-1])[:MAX_SLUG].strip('-') or 'x'
        else:
            part_counter[base] = part_counter.get(base, 0) + 1
            part = 'part-%03d' % part_counter[base]
        ident = '%s_%s' % (base, part)
        if not re.match(r'[A-Za-z_]', ident):
            ident = 'c' + ident
        unique, n = ident, 1
        while unique.lower() in used_lower:
            n += 1
            unique = '%s_dup%d' % (ident, n)      # same code and clause number as an earlier row
        used_lower.add(unique.lower())
        row['id'] = unique
        rows.append(row)

    os.makedirs(args.out, exist_ok=True)
    # Law codes already made from a spreadsheet in the new layout
    # (tools/law-code-from-xlsx.py, e.g. Æthelberht) are better than this
    # spreadsheet's rows for them, so their files and catalogue lines are kept.
    kept = {}
    index_path = os.path.join(args.out, 'index.json')
    if os.path.exists(index_path):
        for c in json.load(open(index_path, encoding='utf-8')):
            p = os.path.join(args.out, c['file'])
            if os.path.exists(p) and 'type="liebermannText"' in open(p, encoding='utf-8').read():
                kept.setdefault(c['code'], []).append(c)
    if kept:
        print('kept (new layout, not replaced): ' + ', '.join('%s (%d clauses)' % (k, len(v)) for k, v in kept.items()))
    kept_files = {c['file'] for v in kept.values() for c in v}
    for name in os.listdir(args.out):
        if (name.endswith('.xml') and name not in kept_files) or name == 'index.json':
            os.remove(os.path.join(args.out, name))

    catalogue = []
    for row in rows:
        if row['code'] in kept:
            catalogue.extend(kept.pop(row['code']))
            continue
        if any(c.get('source') and c['code'] == row['code'] for c in catalogue):
            continue      # a later row of a kept code
        with open(os.path.join(args.out, row['id'] + '.xml'), 'w', encoding='utf-8', newline='\n') as f:
            f.write(build(row, row['block'], row['anchor'], args.date))
        catalogue.append({
            'id': row['id'], 'file': row['id'] + '.xml', 'citation': row['citation'],
            'code': row['code'], 'anchor': row['anchor'], 'clause': row['clause'],
            'citationKeys': [(row['anchor'] or row['code']) + ' ' + k for k in clause_keys(row['clause'])],
            'legalText': row['block'].get('title', ''), 'page': row['page'], 'sheetRow': row['sheetRow'],
        })
    for v in kept.values():      # kept codes this spreadsheet does not have at all
        catalogue.extend(v)
    with open(os.path.join(args.out, 'index.json'), 'w', encoding='utf-8', newline='\n') as f:
        f.write('[\n' + ',\n'.join(json.dumps(c, ensure_ascii=False, separators=(',', ':')) for c in catalogue) + '\n]\n')

    codes = {}
    for row in rows:
        codes.setdefault(row['code'], (row['anchor'], row['anchorHow'], 0))
        a, h, n = codes[row['code']]
        codes[row['code']] = (a, h, n + 1)
    print('clause files written: %d  (sub-header rows read: %d, empty rows skipped: %d)' % (len(rows), headers, skipped_empty))
    print('rows with no clause number: %d' % sum(1 for r in rows if not r['clause']))
    print('law codes: %d, linked to the law code list: %d' % (len(codes), sum(1 for a, h, n in codes.values() if a)))
    for code, (a, how, n) in sorted(codes.items()):
        if how != 'matched':
            print('  %-8s %-34s -> %s  (%d clauses)' % (how, code, a or '(no link)', n))


if __name__ == '__main__':
    sys.exit(main())
