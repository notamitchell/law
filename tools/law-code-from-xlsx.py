#!/usr/bin/env python3
"""
Makes the clause files of ONE law code from a spreadsheet in the new
(2026-10) layout, which gives each clause Liebermann's text, a modern
editor's edition and translation, and a second translation.

    python3 tools/law-code-from-xlsx.py Aethelberht.xlsx
    python3 tools/law-code-from-xlsx.py Aethelberht.xlsx --dry-run

Needs Python 3 and openpyxl (pip install openpyxl).

Only the files of the law code(s) in the spreadsheet are replaced (for
example every oe-laws/Abt_*.xml); every other file in oe-laws/ is left
alone. Their lines in oe-laws/index.json are replaced in place, so the
order of the catalogue and the file names (which the dictionary links to)
stay the same. See docs/OE-LAWS.md, "The new layout".

Columns of the spreadsheet (row 1 is the header row; the column is found by
its header text, so columns may be moved, but not renamed):

  A  code and clause joined (a formula)       M  short name of the modern editor
  B  Liebermann's code abbreviation           N  MS used by the modern editor
  C  Liebermann's clause number               O  modern editor's clause number
  D  full English name of the code            P  modern editor's text
  E  issuing king (and reign)                 Q  modern editor's translation
  F  place and date of promulgation           R  second translation
  G  MSS (Liebermann's sigla; bold = chosen)  S  second translator's name and edition
  H  Liebermann's MS selected for edition     T  supplementary MS
  I  Gesetze page number                      U  translation of the supplementary MS
  J  Liebermann's editorial number            V  offence / legal event
  K  Liebermann's text                        W  people
  L  modern edition (full reference)          X  locations

Italic in the spreadsheet (Liebermann's and the editor's expansions of
abbreviations, and italic words in the translations) is kept as
<hi rend="italic">; italic in a book reference becomes its <title>.
"""
import argparse
import datetime
import importlib.util
import json
import os
import re
import sys
from xml.sax.saxutils import escape

import openpyxl
from openpyxl.cell.rich_text import CellRichText, TextBlock

HERE = os.path.dirname(os.path.abspath(__file__))
_spec = importlib.util.spec_from_file_location('laws_from_xlsx', os.path.join(HERE, 'laws-from-xlsx.py'))
old = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(old)       # shared helpers: slugify, clause keys, law code anchors, witnesses

LIEBERMANN_REF = ('<author>Felix Liebermann</author>, <title>Die Gesetze der Angelsachsen</title>, '
                  'vol. <biblScope unit="volume">1</biblScope> (Halle: Niemeyer, 1903)')

# Header text (start of it, case and spaces ignored) -> our name for the column.
HEADERS = [
    ("liebermann's full code abbreviation", 'citation'),
    ("liebermann's code abbreviation", 'code'),
    ("liebermann's internal clause numbering", 'clause'),
    ('full english name', 'codeName'),
    ('issuing king', 'king'),
    ('place and date of promulgation', 'promulgation'),
    ('mss', 'mss'),
    ("liebermann's ms selected", 'liebermannMS'),
    ('gesetze', 'page'),
    ('editorial cap. numbers of selected ms text, liebermann', 'liebermannNumber'),
    ("liebermann's text", 'liebermannText'),
    ('modern edition', 'editionRef'),
    ('short name of modern editor', 'editor'),
    ('ms used by modern editor', 'editorMS'),
    ('editorial cap. numbers of selected ms text, modern editor', 'editorNumber'),
    ("modern editor's text", 'editionText'),
    ("modern editor's translation", 'translation'),
    ('second translation', 'secondTranslation'),
    ("second translator's name", 'secondRef'),
    ('supplementary ms', 'supplementary'),
    ('translation of supplementary ms', 'supplementaryTranslation'),
    ('offence', 'offence'),
    ('people', 'people'),
    ('locations', 'locations'),
]

ILLEGAL_XML = old.ILLEGAL_XML


def norm_header(text):
    return re.sub(r'\s+', ' ', str(text or '')).strip().lower()


def columns(ws):
    found = {}
    for c in range(1, ws.max_column + 1):
        h = norm_header(ws.cell(1, c).value)
        if not h:
            continue
        for start, name in HEADERS:
            if h.startswith(start) and name not in found:
                found[name] = c
                break
    missing = [n for _, n in HEADERS if n not in found and n not in ('supplementary', 'supplementaryTranslation', 'locations')]
    if missing:
        sys.exit('These columns were not found in row 1: ' + ', '.join(missing))
    return found


def runs(value, base_italic=False):
    """A cell as [(text, italic)], with spaces tidied as in laws-from-xlsx.py."""
    if value is None:
        return []
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    parts = []
    if isinstance(value, CellRichText):
        for p in value:
            if isinstance(p, TextBlock):
                it = p.font.i if p.font is not None and p.font.i is not None else base_italic
                parts.append((p.text, bool(it)))
            else:
                parts.append((str(p), base_italic))
    else:
        parts.append((str(value), base_italic))
    chars = []
    for text, it in parts:
        for ch in ILLEGAL_XML.sub('', text).replace('\xa0', ' '):
            if ch.isspace():
                ch = ' '
                if not chars or chars[-1][0] == ' ':
                    continue
            chars.append((ch, it))
    while chars and chars[-1][0] == ' ':
        chars.pop()
    # A space or punctuation mark is italic only inside an italic phrase (between
    # two italic letters), so "scillingu" + italic "m " gives <hi>m</hi> and a
    # space, but an italic book title stays one piece.
    letters = [i for i, (ch, _) in enumerate(chars) if ch.isalnum()]
    fixed = []
    for i, (ch, it) in enumerate(chars):
        if it and not ch.isalnum():
            before = [j for j in letters if j < i]
            after = [j for j in letters if j > i]
            it = bool(before and after and chars[before[-1]][1] and chars[after[0]][1])
        fixed.append((ch, it))
    out = []
    for ch, it in fixed:
        if out and out[-1][1] == it:
            out[-1] = (out[-1][0] + ch, it)
        else:
            out.append((ch, it))
    return out


def plain(value):
    return ''.join(t for t, _ in runs(value))


def marked(value, tag='<hi rend="italic">%s</hi>'):
    return ''.join((tag % escape(t)) if i else escape(t) for t, i in runs(value))


def surname(ref_runs_text):
    """'F. L. Attenborough, The Laws…' -> 'Attenborough'."""
    name = ref_runs_text.split(',', 1)[0].strip()
    return name.split()[-1] if name else ''


def ident(text):
    return old.slugify(text).lower() or 'source'


def build(row, today, source_name):
    a = row['citation']
    code, clause, anchor = row['code'], row['clause'], row['anchor']
    editor, second = row['editor'], row['secondName']
    ed_id, second_id = ident(editor), ident(second)
    L = []
    w = L.append
    w('<?xml version="1.0" encoding="UTF-8"?>')
    w('<TEI xmlns="http://www.tei-c.org/ns/1.0" xml:id="%s">' % old.attr(row['id']))
    w('  <teiHeader>')
    w('    <fileDesc>')
    w('      <titleStmt>')
    w('        <title>%s</title>' % escape(a))
    w('        <respStmt><resp>Edition of the text</resp><name ref="#liebermann">Liebermann</name></respStmt>')
    if editor and row['editionText']:
        w('        <respStmt><resp>Edition of the text</resp><name ref="#%s">%s</name></respStmt>' % (ed_id, escape(editor)))
    if editor and row['translation']:
        w('        <respStmt><resp>Translation</resp><name ref="#%s">%s</name></respStmt>' % (ed_id, escape(editor)))
    if second and row['secondTranslation']:
        w('        <respStmt><resp>Translation</resp><name ref="#%s">%s</name></respStmt>' % (second_id, escape(second)))
    w('      </titleStmt>')
    w('      <publicationStmt>')
    w('        <publisher>%s</publisher>' % escape(old.PUBLISHER))
    w('        <idno type="slug">%s</idno>' % escape(row['id']))
    w('        <idno type="citation">%s</idno>' % escape(a))
    w('        <idno type="lawCode">%s</idno>' % escape(code))
    if anchor:
        w('        <idno type="lawCodeAnchor">%s</idno>' % escape(anchor))
    if clause:
        w('        <idno type="clause">%s</idno>' % escape(clause))
        for key in old.clause_keys(clause):
            w('        <idno type="citationKey">%s</idno>' % escape((anchor or code) + ' ' + key))
    w('      </publicationStmt>')
    w('      <sourceDesc>')
    w('        <bibl type="legalText"><title>%s</title>%s</bibl>' % (
        escape(row['codeName'] or row['king']),
        (' <note type="reign">%s</note>' % escape(row['king'])) if row['king'] else ''))
    w('        <listBibl type="sources">')
    w('          <bibl xml:id="liebermann" type="edition">%s%s</bibl>' % (
        LIEBERMANN_REF, (', p. <biblScope unit="page">%s</biblScope>' % escape(row['page'])) if row['page'] else ''))
    if editor:
        w('          <bibl xml:id="%s" type="edition">%s</bibl>' % (ed_id, row['editionRefXml']))
    if second and row['secondTranslation']:
        w('          <bibl xml:id="%s" type="translation">%s</bibl>' % (second_id, row['secondRefXml']))
    w('        </listBibl>')
    if row['witnesses']:
        w('        <listWit>')
        for text, selected in row['witnesses']:
            w('          <witness xml:id="wit-%s"%s>%s</witness>' % (ident(text), ' rend="bold"' if selected else '', escape(text)))
        w('        </listWit>')
    w('      </sourceDesc>')
    w('    </fileDesc>')
    w('    <encodingDesc>')
    w(old.PREFIX_DEFS)
    w('    </encodingDesc>')
    kws = [('offence-legal-event', row['offence']), ('people', row['people']), ('locations', row['locations'])]
    w('    <profileDesc>')
    if row['promulgation']:
        w('      <creation><note type="promulgation">%s</note></creation>' % escape(row['promulgation']))
    if any(t for _, t in kws):
        w('      <textClass>')
        for name, terms in kws:
            if terms:
                w('        <keywords n="%s">%s</keywords>' % (name, ''.join('<term>%s</term>' % escape(t) for t in old.split_terms(terms))))
        w('      </textClass>')
    w('    </profileDesc>')
    w('    <revisionDesc>')
    w('      <change when="%s">Generated from row %d of %s.</change>' % (today, row['sheetRow'], escape(source_name)))
    w('    </revisionDesc>')
    w('  </teiHeader>')
    w('  <text>')
    w('    <body>')
    w('      <div type="clause"%s>' % ((' n="%s"' % old.attr(clause)) if clause else ''))
    head_code = '<ref type="lawCode" target="law:%s">%s</ref>' % (old.attr(anchor), escape(code)) if anchor else escape(code)
    w('        <head>%s%s</head>' % (head_code, (' ' + escape(clause)) if clause else ''))

    def ab(kind, lang, source, xml, n='', ms=''):
        if not xml:
            return
        at = ' type="%s" xml:lang="%s" source="#%s"' % (kind, lang, source)
        if n:
            at += ' n="%s"' % old.attr(n)
        if ms:
            at += ' corresp="#wit-%s"' % ident(ms)
        w('        <ab%s>%s</ab>' % (at, xml))
    ab('liebermannText', 'ang', 'liebermann', row['liebermannTextXml'], row['liebermannNumber'], row['liebermannMS'])
    ab('editionText', 'ang', ed_id, row['editionTextXml'], row['editorNumber'], row['editorMS'])
    ab('translation', 'en', ed_id, row['translationXml'])
    ab('translation', 'en', second_id, row['secondTranslationXml'])
    if row['supplementaryXml']:
        w('        <ab type="supplementaryText" xml:lang="ang">%s</ab>' % row['supplementaryXml'])
    if row['supplementaryTranslationXml']:
        w('        <ab type="supplementaryTranslation" xml:lang="en">%s</ab>' % row['supplementaryTranslationXml'])
    w('      </div>')
    w('    </body>')
    w('  </text>')
    w('</TEI>')
    return '\n'.join(L) + '\n'


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('xlsx')
    ap.add_argument('--out', default='oe-laws')
    ap.add_argument('--mark-law', default='data/markLaw.json')
    ap.add_argument('--date', default=datetime.date.today().isoformat())
    ap.add_argument('--dry-run', action='store_true', help='report only; write nothing')
    args = ap.parse_args()

    ws = openpyxl.load_workbook(args.xlsx, rich_text=True).active
    col = columns(ws)
    find_anchor = old.anchor_finder(args.mark_law)
    source_name = 'the spreadsheet ' + os.path.basename(args.xlsx)

    rows = []
    for r in range(2, ws.max_row + 1):
        cell = lambda name: ws.cell(r, col[name]) if name in col else None
        get = lambda name: old.clean(plain(cell(name).value)) if cell(name) is not None else ''
        xml = lambda name: marked(cell(name).value) if cell(name) is not None else ''
        code = get('code')
        if not code:
            continue
        clause = get('clause')
        anchor, how = find_anchor(code)
        edition_ref = cell('editionRef').value
        second_ref = cell('secondRef').value
        row = {
            'sheetRow': r, 'code': code, 'clause': clause, 'anchor': anchor,
            'citation': ('%s %s' % (code, clause)).strip(),
            'codeName': get('codeName'), 'king': get('king'), 'promulgation': get('promulgation'),
            'page': get('page'), 'liebermannMS': get('liebermannMS'), 'liebermannNumber': get('liebermannNumber'),
            'editor': get('editor'), 'editorMS': get('editorMS'), 'editorNumber': get('editorNumber'),
            'liebermannTextXml': xml('liebermannText'), 'editionTextXml': xml('editionText'),
            'editionText': get('editionText'),
            'translationXml': xml('translation'), 'translation': get('translation'),
            'secondTranslationXml': xml('secondTranslation'), 'secondTranslation': get('secondTranslation'),
            'supplementaryXml': xml('supplementary'), 'supplementaryTranslationXml': xml('supplementaryTranslation'),
            'editionRefXml': marked(edition_ref, '<title>%s</title>'),
            'secondRefXml': marked(second_ref, '<title>%s</title>'),
            'secondName': surname(plain(second_ref)),
            'offence': get('offence'), 'people': get('people'), 'locations': get('locations'),
            'witnesses': old.witnesses(cell('mss')),
        }
        base = anchor or old.slugify(code) or 'code'
        parts = old.clause_parts(clause)
        part = old.slugify(parts[0] if len(parts) == 1 else parts[0] + ' to ' + parts[-1])[:old.MAX_SLUG].strip('-') or 'x'
        row['id'] = '%s_%s' % (base, part)
        rows.append(row)

    ids = [r['id'] for r in rows]
    dups = {i for i in ids if ids.count(i) > 1}
    if dups:
        sys.exit('Two rows would make the same file: ' + ', '.join(sorted(dups)))
    codes = sorted({r['code'] for r in rows})

    index_path = os.path.join(args.out, 'index.json')
    catalogue = json.load(open(index_path, encoding='utf-8'))
    old_ids = {c['id'] for c in catalogue if c['code'] in codes}
    new_ids = set(ids)
    print('law code(s): %s; clause rows: %d' % (', '.join(codes), len(rows)))
    print('files replaced: %d; new: %d; old files with no row now (removed): %d' % (
        len(old_ids & new_ids), len(new_ids - old_ids), len(old_ids - new_ids)))
    for i in sorted(new_ids - old_ids):
        print('  new     ', i)
    for i in sorted(old_ids - new_ids):
        print('  removed ', i)
    if args.dry_run:
        return

    for i in old_ids - new_ids:
        p = os.path.join(args.out, i + '.xml')
        if os.path.exists(p):
            os.remove(p)
    for row in rows:
        with open(os.path.join(args.out, row['id'] + '.xml'), 'w', encoding='utf-8', newline='\n') as f:
            f.write(build(row, args.date, source_name))

    def entry(row):
        return {
            'id': row['id'], 'file': row['id'] + '.xml', 'citation': row['citation'],
            'code': row['code'], 'anchor': row['anchor'], 'clause': row['clause'],
            'citationKeys': [(row['anchor'] or row['code']) + ' ' + k for k in old.clause_keys(row['clause'])],
            'legalText': row['king'].split(',')[0].strip() or row['codeName'], 'page': row['page'], 'sheetRow': row['sheetRow'],
            'source': os.path.basename(args.xlsx),
        }
    # Put the code's clauses where its first clause was, in spreadsheet order.
    out, placed = [], False
    for c in catalogue:
        if c['code'] in codes:
            if not placed:
                out.extend(entry(r) for r in rows)
                placed = True
            continue
        out.append(c)
    if not placed:
        out.extend(entry(r) for r in rows)
    with open(index_path, 'w', encoding='utf-8', newline='\n') as f:
        f.write('[\n' + ',\n'.join(json.dumps(c, ensure_ascii=False, separators=(',', ':')) for c in out) + '\n]\n')
    print('written to %s/ and %s' % (args.out, index_path))


if __name__ == '__main__':
    sys.exit(main())
